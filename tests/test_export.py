"""genMOLDB 로 나가는 CSV 계약을 고정하는 테스트.

molgen/CSV_COLUMNS.md §0: "로더가 비어 있지 않은 동명 컬럼을 순서대로 모아
하나의 리스트로 만든다" — 즉 자식 행의 순서는 의미를 갖는다.
"""

import csv

from openpyxl import load_workbook

from genviewer.service.repo.io import IORepository
from genviewer.utils.exporter import write_csv, write_xlsx

HEADERS = ["CellType", "Param"]
ROWS = [["=1+1", "x"], ["a=b", "=SUM(A1:A2)"], ["plain", ""]]


def _xlsx_cells(path):
    """다시 연 xlsx 의 셀을 행마다 (값, data_type) 로 돌려준다."""
    ws = load_workbook(path).active
    return [[(cell.value, cell.data_type) for cell in row] for row in ws.iter_rows()]


def test_xlsx_writes_values_starting_with_equals_as_text(tmp_path):
    """'=' 로 시작하는 값은 수식이 아니라 글자 그대로의 문자열 셀이다(K12, 2026-10-17 승인).

    openpyxl 은 '=' 로 시작하는 문자열을 수식(data_type "f")으로 저장해서, 엑셀이
    DB 의 값 대신 계산 결과를 보여 줬다. 값 앞에 ' 를 붙이는 것도 아니다.
    """
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), HEADERS, ROWS)

    cells = _xlsx_cells(path)
    assert cells[1][0] == ("=1+1", "s")
    assert cells[2][1] == ("=SUM(A1:A2)", "s")


def test_xlsx_writes_other_values_as_before(tmp_path):
    """헤더와 '=' 로 시작하지 않는 값은 그대로다. 빈 문자열은 글자 없는 셀이다."""
    path = tmp_path / "out.xlsx"
    write_xlsx(str(path), HEADERS, ROWS)

    cells = _xlsx_cells(path)
    assert len(cells) == 4
    assert cells[0] == [("CellType", "s"), ("Param", "s")]
    assert cells[1][1] == ("x", "s")
    assert cells[2][0] == ("a=b", "s")
    assert cells[3] == [("plain", "s"), (None, "inlineStr")]


def test_csv_keeps_values_starting_with_equals_as_written(tmp_path):
    """CSV 는 원문 그대로다. molgen 이 이 글자를 읽는다."""
    path = tmp_path / "out.csv"
    write_csv(str(path), HEADERS, ROWS)

    with open(path, newline="", encoding="utf-8") as f:
        assert list(csv.reader(f)) == [
            ["CellType", "Param"],
            ["=1+1", "x"],
            ["a=b", "=SUM(A1:A2)"],
            ["plain", ""],
        ]


def test_each_row_carries_its_own_process_name(db, two_processes):
    """device 마다 CdsLibName 은 자기 process 여야 한다.

    회귀 방지: fetch_export_data 의 2패스가 1패스에서 새어나온 process_name
    변수를 읽어서, 모든 행이 마지막 device 의 process 이름을 갖고 있었다.
    """
    headers, rows = IORepository(db).fetch_export_data(two_processes)
    cds = headers.index("CdsLibName")
    inst = headers.index("InstCellName")
    got = {row[inst]: row[cds] for row in rows}
    assert got == {"nfet_b": "sf2p", "pfet_b": "sf4u"}


def test_description_follows_inst_cell_name(db, two_processes):
    """Description 은 InstCellName 다음 칸이고 값은 device_info 의 description 이다(K44)."""
    info = [{**two_processes[0], "description": "ESD 용"}, two_processes[1]]
    headers, rows = IORepository(db).fetch_export_data(info)

    assert headers[:4] == ["CellType", "CdsLibName", "InstCellName", "Description"]
    assert {row[2]: row[3] for row in rows} == {"nfet_b": "ESD 용", "pfet_b": ""}


def test_child_row_order_is_preserved(db, two_processes):
    """InstTerm 컬럼들은 삽입 순서대로 나와야 한다.

    회귀 방지: 자식 테이블 조회에 ORDER BY 가 없었다. SQLite 는 rowid 순서를
    돌려주지만 PostgreSQL 은 heap 순서라서 UPDATE 후 뒤섞인다.
    """
    headers, rows = IORepository(db).fetch_export_data(two_processes)
    inst = headers.index("InstCellName")
    term_cols = [i for i, h in enumerate(headers) if h == "InstTerm"]
    assert len(term_cols) == 4, "device 당 InstTerm 행이 4개이므로 컬럼도 4개여야 한다"

    for row in rows:
        assert [row[i] for i in term_cols] == ["DRN d", "GATE g", "SRC s", "BULK b"], (
            f"{row[inst]} 의 InstTerm 순서가 깨졌다"
        )

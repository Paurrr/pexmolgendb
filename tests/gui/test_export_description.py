"""device 표 내보내기는 화면의 description 을 Description 열에 쓴다 (phase 18 step 2, K44).

사용자 승인 2026-10-17: "내보내기(GUI·CLI, CSV·xlsx)는 InstCellName 다음에 Description
컬럼을 쓴다". 값은 name·celltype 처럼 화면(모델)의 값이다. 저장하지 않은 편집이면 그 편집 값이고,
DB 에서 다시 읽지 않는다. 쉼표·줄바꿈을 바꿔 쓰는 것은 step 3(CSV 쓰기)이고, 그 GUI 경로는
마지막 두 테스트가 본다.
"""

import csv
from pathlib import Path

from gui_support.helpers import edit_device_cell
from openpyxl import load_workbook


def export_to(ctrl, path: Path, fmt: str) -> None:
    """저장 경로 대화상자 대신 path 를 돌려주고 fmt 메뉴로 내보낸다."""
    ctrl.device_table.ui._ask_export_path = lambda _suggested: str(path)
    ctrl.device_table.export_selected_devices(fmt)


def descriptions_in_csv(path: Path) -> dict[str, str]:
    """InstCellName → Description. 헤더 4번째가 Description 이어야 한다."""
    with open(path, newline="", encoding="utf-8") as f:
        headers, *rows = list(csv.reader(f))
    assert headers[:4] == ["CellType", "CdsLibName", "InstCellName", "Description"]
    return {row[2]: row[3] for row in rows}


def descriptions_in_xlsx(path: Path) -> dict[str, str]:
    ws = load_workbook(path).active
    headers, *rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert headers[:4] == ["CellType", "CdsLibName", "InstCellName", "Description"]
    # openpyxl 은 빈 문자열 셀을 None 으로 읽는다.
    return {row[2]: row[3] or "" for row in rows}


def test_csv_export_writes_the_saved_description(device_tab_ctrl, store, dialogs, tmp_path):
    ctrl = device_tab_ctrl
    store.touch_device(store.find_device("nfet_a")["id"], actor="kim", description="ESD 용")
    ctrl.device_table.refresh_devices()
    path = tmp_path / "out.csv"

    export_to(ctrl, path, "csv")

    assert descriptions_in_csv(path) == {"nfet_a": "ESD 용", "nfet_b": ""}
    assert dialogs.shown == [("info", "Save CSV", "Success!\n - out.csv")]


def test_export_writes_the_unsaved_description_on_screen(
    device_tab_ctrl, store, dialogs, tmp_path
):
    ctrl = device_tab_ctrl
    assert ctrl.device_table.ui.model_safe().get_device(1)["name"] == "nfet_b"
    edit_device_cell(ctrl, 1, "description", "화면 값")
    path = tmp_path / "out.csv"

    export_to(ctrl, path, "csv")

    assert descriptions_in_csv(path) == {"nfet_a": "", "nfet_b": "화면 값"}
    assert store.find_device("nfet_b")["description"] == ""  # 저장하지 않았다
    assert dialogs.shown == [("info", "Save CSV", "Success!\n - out.csv")]


def test_xlsx_export_writes_the_description(device_tab_ctrl, store, dialogs, tmp_path):
    ctrl = device_tab_ctrl
    store.touch_device(store.find_device("nfet_a")["id"], actor="kim", description="ESD 용")
    ctrl.device_table.refresh_devices()
    edit_device_cell(ctrl, 1, "description", "화면 값")
    path = tmp_path / "out.xlsx"

    export_to(ctrl, path, "xlsx")

    assert descriptions_in_xlsx(path) == {"nfet_a": "ESD 용", "nfet_b": "화면 값"}
    assert dialogs.shown == [("info", "Save XLSX", "Success!\n - out.xlsx")]


def test_csv_export_turns_description_commas_into_semicolons(
    device_tab_ctrl, store, dialogs, tmp_path
):
    """molgen 은 CSV 한 줄을 ',' 로만 나눈다. 설명의 쉼표가 뒤 컬럼을 밀지 않는다.

    사용자 승인 2026-10-18. DB 의 description 은 그대로다.
    """
    ctrl = device_tab_ctrl
    store.touch_device(store.find_device("nfet_a")["id"], actor="kim", description="ESD 용, 2단")
    ctrl.device_table.refresh_devices()
    path = tmp_path / "out.csv"

    export_to(ctrl, path, "csv")

    assert descriptions_in_csv(path)["nfet_a"] == "ESD 용; 2단"
    header, *lines = path.read_text(encoding="utf-8").splitlines()
    (nfet_a,) = [line for line in lines if line.split(",")[2] == "nfet_a"]
    assert len(nfet_a.split(",")) == len(header.split(","))
    assert store.find_device("nfet_a")["description"] == "ESD 용, 2단"  # DB 값은 그대로
    assert dialogs.shown == [("info", "Save CSV", "Success!\n - out.csv")]


def test_xlsx_export_keeps_description_commas(device_tab_ctrl, store, dialogs, tmp_path):
    ctrl = device_tab_ctrl
    store.touch_device(store.find_device("nfet_a")["id"], actor="kim", description="ESD 용, 2단")
    ctrl.device_table.refresh_devices()
    path = tmp_path / "out.xlsx"

    export_to(ctrl, path, "xlsx")

    assert descriptions_in_xlsx(path)["nfet_a"] == "ESD 용, 2단"
    assert dialogs.shown == [("info", "Save XLSX", "Success!\n - out.xlsx")]
